"""US5 test: Excel export parity (FR-021, SC-010)."""

from __future__ import annotations

from decimal import Decimal

from openpyxl import load_workbook

from export.excel import FIRST_DATA_ROW, FIRST_VALUE_COL, column_headers, write_pnl
from tests.conftest import cells_by_key


def test_excel_preserves_structure_and_figures(result6, cfg6, tmp_path):
    out = tmp_path / "pnl.xlsx"
    write_pnl(result6, out)
    wb = load_workbook(out)
    ws = wb["P&L"]

    # Line order preserved in column A.
    written_lines = [
        ws.cell(row=FIRST_DATA_ROW + i, column=1).value for i in range(len(result6.lines))
    ]
    assert written_lines == [line.reporting_line for line in result6.lines]

    # Column headers match scenario x time-cut order.
    headers = [ws.cell(row=1, column=FIRST_VALUE_COL + i).value for i in range(3 * 4)]
    assert headers == column_headers()

    # Every figure equals the in-memory result at configured precision.
    mp, rp = cfg6.settings.money_precision, cfg6.settings.ratio_precision
    for r_off, line in enumerate(result6.lines):
        for c_off, cell in enumerate(line.cells):
            read = ws.cell(row=FIRST_DATA_ROW + r_off, column=FIRST_VALUE_COL + c_off).value
            precision = rp if cell.is_margin else mp
            got = Decimal(str(read)).quantize(Decimal(1).scaleb(-precision))
            assert got == cell.value, f"{line.reporting_line} cell mismatch: {got} != {cell.value}"
